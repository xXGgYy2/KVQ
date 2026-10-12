#!/usr/bin/env python3
"""Evaluate KVQ branches on 4K-HEVC and VDPVE datasets.

Run this script from the KVQ repository root. It writes:
  - parsed_labels_4k.csv / parsed_labels_vdpve.csv
  - per_video.csv
  - metrics.csv
  - kvq_results.xlsx (when openpyxl is installed)
"""

from __future__ import annotations

import argparse
import csv
import inspect
import math
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import torch
import yaml
from scipy.optimize import curve_fit
from scipy.stats import kendalltau, pearsonr, spearmanr
from torch.utils.data import DataLoader
from tqdm import tqdm

import vqa.datasets as kvq_datasets
import vqa.models as kvq_models


VIDEO_EXTENSIONS = {
    ".avi",
    ".m4v",
    ".mkv",
    ".mov",
    ".mp4",
    ".mpeg",
    ".mpg",
    ".webm",
    ".y4m",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Evaluate every scalar KVQ output branch and compute VQA correlations."
    )
    parser.add_argument("--config", default="configs/single_video.yaml")
    parser.add_argument("--checkpoint", default="weights/KVQ.pth")
    parser.add_argument(
        "--dataset1-dir",
        default=(
            "/mnt/data/local-disk2/xyy/datasets/"
            "4K-HEVC-Video-Subjective-Quality-Database/mp4"
        ),
    )
    parser.add_argument(
        "--dataset1-pdf",
        default=(
            "/mnt/data/local-disk2/xyy/datasets/"
            "4K-HEVC-Video-Subjective-Quality-Database/SubScore2.pdf"
        ),
    )
    parser.add_argument(
        "--dataset1-label-csv",
        default=None,
        help="Optional verified CSV with filename,mos columns; overrides PDF parsing.",
    )
    parser.add_argument(
        "--dataset2-dir",
        default="/mnt/data/local-disk2/xyy/datasets/VDPVE/train",
    )
    parser.add_argument(
        "--dataset2-labels",
        default="/mnt/data/local-disk2/xyy/datasets/VDPVE/train.txt",
    )
    parser.add_argument("--output-dir", default="results/kvq_two_datasets")
    parser.add_argument("--device", default="cuda:0")
    parser.add_argument("--num-workers", type=int, default=0)
    parser.add_argument("--no-amp", action="store_true")
    parser.add_argument(
        "--prepare-only",
        action="store_true",
        help="Only parse labels and write audit CSVs; do not run KVQ.",
    )
    parser.add_argument(
        "--allow-missing-labels",
        action="store_true",
        help="Skip videos whose MOS cannot be parsed instead of failing.",
    )
    return parser.parse_args()


def find_videos(root: Path) -> list[Path]:
    if not root.is_dir():
        raise FileNotFoundError(f"Video directory does not exist: {root}")
    videos = sorted(
        path for path in root.rglob("*") if path.is_file() and path.suffix.lower() in VIDEO_EXTENSIONS
    )
    if not videos:
        raise RuntimeError(f"No video files found under: {root}")
    return videos


def numeric_tokens(text: str) -> list[float]:
    return [
        float(value)
        for value in re.findall(r"(?<![A-Za-z0-9_.])[-+]?\d+(?:\.\d+)?", text)
    ]


def read_verified_csv(path: Path, videos: list[Path], root: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError(f"CSV has no header: {path}")
        normalized = {name.strip().lower(): name for name in reader.fieldnames}
        filename_key = normalized.get("filename") or normalized.get("video") or normalized.get("name")
        mos_key = normalized.get("mos") or normalized.get("score") or normalized.get("dmos")
        if not filename_key or not mos_key:
            raise ValueError(f"CSV must contain filename/video and mos/score columns: {path}")
        scores = {
            Path(row[filename_key].strip()).name.lower(): float(row[mos_key])
            for row in reader
            if row.get(filename_key) and row.get(mos_key)
        }

    rows = []
    for video in videos:
        key = video.name.lower()
        rows.append(
            {
                "filename": video.relative_to(root).as_posix(),
                "mos": scores.get(key),
                "source": str(path),
                "matched_text": key,
            }
        )
    return rows


def extract_pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError(
            "PDF parsing requires pypdf. Install it with: python -m pip install pypdf"
        ) from exc

    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages)
    if not text.strip():
        raise RuntimeError(f"No text could be extracted from PDF: {path}")
    return text


def canonical_4k_video_id(name: str) -> str:
    """Map Campfire_Party_R1.mp4 and Campfire_Party1 to the same key."""
    stem = Path(name).stem
    stem = re.sub(r"_R(?=\d+$)", "", stem, flags=re.IGNORECASE)
    return re.sub(r"[^a-z0-9]", "", stem.lower())


def parse_pdf_labels(pdf: Path, videos: list[Path], root: Path) -> list[dict[str, Any]]:
    text = extract_pdf_text(pdf)
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]
    table: dict[str, tuple[float, str]] = {}
    number = r"[-+]?\d+(?:\.\d+)?"
    row_pattern = re.compile(
        rf"^(?P<name>\S+)\s+(?P<bitrate>{number})\s+"
        rf"(?P<mos>{number})\s+(?P<ci>{number})$"
    )
    for line in lines:
        match = row_pattern.match(line)
        if not match:
            continue
        table[canonical_4k_video_id(match.group("name"))] = (
            float(match.group("mos")),
            line,
        )

    rows = []
    for video in videos:
        score_and_line = table.get(canonical_4k_video_id(video.name))
        score, source_line = score_and_line if score_and_line else (None, "")
        rows.append(
            {
                "filename": video.relative_to(root).as_posix(),
                "mos": score,
                "source": str(pdf),
                "matched_text": source_line,
            }
        )
    return rows


def parse_text_labels(label_file: Path, videos: list[Path], root: Path) -> list[dict[str, Any]]:
    lines = label_file.read_text(encoding="utf-8-sig", errors="replace").splitlines()
    rows = []
    for video in videos:
        score = None
        source_line = ""
        for line in lines:
            lowered = line.lower()
            token = video.name.lower() if video.name.lower() in lowered else video.stem.lower()
            if token not in lowered:
                continue
            cleaned = re.sub(re.escape(video.name), " ", line, flags=re.IGNORECASE)
            cleaned = re.sub(re.escape(video.stem), " ", cleaned, flags=re.IGNORECASE)
            values = numeric_tokens(cleaned)
            if values:
                score = values[-1]
                source_line = line.strip()
                break
        rows.append(
            {
                "filename": video.relative_to(root).as_posix(),
                "mos": score,
                "source": str(label_file),
                "matched_text": source_line,
            }
        )
    return rows


def write_dict_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fieldnames is None:
        fieldnames = list(rows[0]) if rows else []
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def validate_labels(
    dataset: str,
    rows: list[dict[str, Any]],
    allow_missing: bool,
) -> list[dict[str, Any]]:
    missing = [row["filename"] for row in rows if row["mos"] is None]
    print(f"[{dataset}] videos={len(rows)}, labels={len(rows) - len(missing)}, missing={len(missing)}")
    if missing:
        print(f"[{dataset}] first missing labels: {missing[:10]}")
        if not allow_missing:
            raise RuntimeError(
                f"{dataset}: {len(missing)} MOS values were not parsed. "
                "Inspect the parsed_labels CSV, then provide a verified CSV or use "
                "--allow-missing-labels."
            )
    return [row for row in rows if row["mos"] is not None]


def create_kvq_annotation(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        for row in rows:
            handle.write(f"{row['filename']}, -1, -1, {float(row['mos']):.12g}\n")


def instantiate_dataset(dataset_type: str, options: dict[str, Any]):
    dataset_class = getattr(kvq_datasets, dataset_type)
    try:
        return dataset_class(options)
    except TypeError as positional_error:
        try:
            return dataset_class(**options)
        except TypeError:
            raise positional_error


def instantiate_model(model_type: str, options: dict[str, Any]):
    model_class = getattr(kvq_models, model_type)
    try:
        return model_class(**options)
    except TypeError as keyword_error:
        try:
            return model_class(options)
        except TypeError:
            raise keyword_error


def load_model(config: dict[str, Any], checkpoint: Path, device: torch.device):
    model_options = config["model"]
    model = instantiate_model(model_options["type"], model_options.get("args", {}))
    try:
        checkpoint_data = torch.load(checkpoint, map_location="cpu", weights_only=False)
    except TypeError:
        checkpoint_data = torch.load(checkpoint, map_location="cpu")
    state_dict = checkpoint_data.get("state_dict", checkpoint_data)
    try:
        model.load_state_dict(state_dict, strict=True)
    except RuntimeError:
        stripped = {
            key.removeprefix("module."): value for key, value in state_dict.items()
        }
        model.load_state_dict(stripped, strict=True)
    model.to(device).eval()
    return model


def reduce_tensor(value: torch.Tensor) -> float:
    return float(value.detach().float().mean().cpu().item())


def flatten_outputs(value: Any, prefix: str, output: dict[str, float]) -> None:
    if torch.is_tensor(value):
        output[prefix] = reduce_tensor(value)
    elif isinstance(value, dict):
        for key, child in value.items():
            flatten_outputs(child, f"{prefix}.{key}", output)
    elif isinstance(value, (list, tuple)):
        for index, child in enumerate(value):
            flatten_outputs(child, f"{prefix}.{index}", output)
    elif isinstance(value, (float, int, np.number)):
        output[prefix] = float(value)


class BranchCapture:
    """Capture the scalar mean from KVQ scoring/fusion modules."""

    TARGETS = {"vqa_head", "weight_head", "weight_conv"}

    def __init__(self, model: torch.nn.Module):
        self.values: dict[str, list[float]] = defaultdict(list)
        self.handles = []
        for name, module in model.named_modules():
            if name.split(".")[-1] in self.TARGETS:
                self.handles.append(module.register_forward_hook(self._hook(name)))

    def _hook(self, name: str):
        def save(_module, _inputs, value):
            flattened: dict[str, float] = {}
            flatten_outputs(value, name, flattened)
            for key, score in flattened.items():
                self.values[key].append(score)

        return save

    def reset(self) -> None:
        self.values.clear()

    def results(self) -> dict[str, float]:
        results = {}
        for key, values in self.values.items():
            if len(values) == 1:
                results[f"hook.{key}"] = values[0]
            else:
                for index, score in enumerate(values):
                    results[f"hook.{key}.{index}"] = score
        return results

    def close(self) -> None:
        for handle in self.handles:
            handle.remove()


def model_forward(model: torch.nn.Module, inputs: dict[str, torch.Tensor]):
    signature = inspect.signature(model.forward)
    kwargs = {}
    defaults = {
        "cross_flag": False,
        "epoch": 0,
        "inference": True,
    }
    for name, value in defaults.items():
        if name in signature.parameters:
            kwargs[name] = value
    return model(inputs, **kwargs)


def build_dataset(
    config: dict[str, Any],
    annotation: Path,
    data_root: Path,
):
    data_templates = config.get("data", {})
    template = next(
        (value for key, value in data_templates.items() if key.startswith("val")),
        next(iter(data_templates.values()), None),
    )
    if template is None:
        raise ValueError("Config contains no data template.")
    options = dict(template["args"])
    options["phase"] = "test"
    options["anno_file"] = str(annotation.resolve())
    options["data_prefix"] = str(data_root.resolve())
    return instantiate_dataset(template["type"], options), options["sample_types"]


def get_batch_name(batch: dict[str, Any], fallback: str) -> str:
    for key in ("name", "filename", "video_name"):
        value = batch.get(key)
        if isinstance(value, (list, tuple)) and value:
            return str(value[0])
        if isinstance(value, str):
            return value
    return fallback


def evaluate_dataset(
    dataset_name: str,
    data_root: Path,
    label_rows: list[dict[str, Any]],
    annotation: Path,
    config: dict[str, Any],
    model: torch.nn.Module,
    device: torch.device,
    num_workers: int,
    amp: bool,
) -> list[dict[str, Any]]:
    dataset, sample_types = build_dataset(config, annotation, data_root)
    loader = DataLoader(
        dataset,
        batch_size=1,
        shuffle=False,
        num_workers=num_workers,
        pin_memory=device.type == "cuda",
    )
    mos_by_name = {
        Path(row["filename"]).name.lower(): float(row["mos"]) for row in label_rows
    }
    capture = BranchCapture(model)
    output_rows = []
    try:
        for index, batch in enumerate(tqdm(loader, desc=dataset_name)):
            inputs = {}
            video_batch = batch.get("video", batch)
            for view_name in sample_types:
                if view_name in video_batch and torch.is_tensor(video_batch[view_name]):
                    inputs[view_name] = video_batch[view_name].to(
                        device, non_blocking=True
                    )
            if not inputs:
                raise RuntimeError(
                    f"No configured video tensors found in batch. Keys: {list(batch)}"
                )

            fallback = label_rows[index]["filename"] if index < len(label_rows) else str(index)
            filename = get_batch_name(batch, fallback)
            key = Path(filename).name.lower()
            mos = mos_by_name.get(key)
            if mos is None and index < len(label_rows):
                mos = float(label_rows[index]["mos"])

            capture.reset()
            with torch.inference_mode():
                with torch.cuda.amp.autocast(enabled=amp and device.type == "cuda"):
                    raw_output = model_forward(model, inputs)

            branches: dict[str, float] = {}
            flatten_outputs(raw_output, "output", branches)
            branches.update(capture.results())
            if not branches:
                raise RuntimeError("KVQ returned no numeric/tensor branch outputs.")

            row: dict[str, Any] = {
                "dataset": dataset_name,
                "filename": filename,
                "mos": mos,
            }
            row.update(branches)
            output_rows.append(row)
            del batch, inputs, raw_output
            if device.type == "cuda":
                torch.cuda.empty_cache()
    finally:
        capture.close()
    return output_rows


def five_parameter_logistic(x, beta1, beta2, beta3, beta4, beta5):
    z = np.clip(beta2 * (x - beta3), -60.0, 60.0)
    return beta1 * (0.5 - 1.0 / (1.0 + np.exp(z))) + beta4 * x + beta5


def safe_stat(function, x: np.ndarray, y: np.ndarray) -> float:
    if len(x) < 2 or np.ptp(x) == 0 or np.ptp(y) == 0:
        return math.nan
    try:
        result = function(x, y)
        return float(result.statistic if hasattr(result, "statistic") else result[0])
    except Exception:
        return math.nan


def logistic_plcc(predictions: np.ndarray, mos: np.ndarray) -> tuple[float, bool]:
    if len(predictions) < 6 or np.ptp(predictions) == 0 or np.ptp(mos) == 0:
        return math.nan, False
    initial = [
        float(np.ptp(mos)),
        1.0 / max(float(np.std(predictions)), 1e-6),
        float(np.mean(predictions)),
        0.0,
        float(np.mean(mos)),
    ]
    try:
        params, _ = curve_fit(
            five_parameter_logistic,
            predictions,
            mos,
            p0=initial,
            maxfev=100000,
        )
        mapped = five_parameter_logistic(predictions, *params)
        return safe_stat(pearsonr, mapped, mos), True
    except Exception:
        return math.nan, False


def compute_metrics(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    datasets = sorted({str(row["dataset"]) for row in rows})
    reserved = {"dataset", "filename", "mos"}
    metrics = []
    for dataset in datasets:
        dataset_rows = [row for row in rows if row["dataset"] == dataset]
        branches = sorted(set().union(*(row.keys() for row in dataset_rows)) - reserved)
        for branch in branches:
            pairs = [
                (float(row[branch]), float(row["mos"]))
                for row in dataset_rows
                if branch in row
                and row[branch] is not None
                and row["mos"] is not None
                and np.isfinite(float(row[branch]))
                and np.isfinite(float(row["mos"]))
            ]
            if not pairs:
                continue
            predictions = np.asarray([pair[0] for pair in pairs], dtype=np.float64)
            mos = np.asarray([pair[1] for pair in pairs], dtype=np.float64)
            mapped_plcc, fitted = logistic_plcc(predictions, mos)
            metrics.append(
                {
                    "dataset": dataset,
                    "branch": branch,
                    "n": len(pairs),
                    "SRCC": safe_stat(spearmanr, predictions, mos),
                    "KRCC": safe_stat(kendalltau, predictions, mos),
                    "PLCC": safe_stat(pearsonr, predictions, mos),
                    "PLCC_logistic": mapped_plcc,
                    # Keep the spelling used in the request as a compatibility alias.
                    "PLCC_logtistic": mapped_plcc,
                    "logistic_fit_ok": fitted,
                }
            )
    return metrics


def rectangularize(rows: list[dict[str, Any]], leading: Iterable[str]) -> tuple[list[str], list[dict[str, Any]]]:
    leading = list(leading)
    all_fields = set().union(*(row.keys() for row in rows)) if rows else set(leading)
    fields = leading + sorted(all_fields - set(leading))
    return fields, [{field: row.get(field, "") for field in fields} for row in rows]


def write_xlsx(path: Path, per_video: list[dict[str, Any]], metrics: list[dict[str, Any]]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[WARN] openpyxl is not installed; CSV files were still written.")
        return False

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows, leading in (
        ("per_video", per_video, ("dataset", "filename", "mos")),
        ("metrics", metrics, ("dataset", "branch", "n")),
    ):
        sheet = workbook.create_sheet(title)
        fields, normalized = rectangularize(rows, leading)
        sheet.append(fields)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in normalized:
            sheet.append([row[field] for field in fields])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column, field in enumerate(fields, 1):
            width = min(max(len(field) + 2, 12), 45)
            sheet.column_dimensions[get_column_letter(column)].width = width
    workbook.save(path)
    return True


def main() -> int:
    args = parse_args()
    repo_root = Path.cwd()
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    dataset1_root = Path(args.dataset1_dir)
    dataset2_root = Path(args.dataset2_dir)
    dataset1_videos = find_videos(dataset1_root)
    dataset2_videos = find_videos(dataset2_root)

    if args.dataset1_label_csv:
        dataset1_rows = read_verified_csv(
            Path(args.dataset1_label_csv), dataset1_videos, dataset1_root
        )
    else:
        dataset1_rows = parse_pdf_labels(
            Path(args.dataset1_pdf), dataset1_videos, dataset1_root
        )
    dataset2_rows = parse_text_labels(
        Path(args.dataset2_labels), dataset2_videos, dataset2_root
    )

    write_dict_csv(output_dir / "parsed_labels_4k.csv", dataset1_rows)
    write_dict_csv(output_dir / "parsed_labels_vdpve.csv", dataset2_rows)
    dataset1_rows = validate_labels("4K_HEVC", dataset1_rows, args.allow_missing_labels)
    dataset2_rows = validate_labels("VDPVE_train", dataset2_rows, args.allow_missing_labels)

    if args.prepare_only:
        print(f"Label audit files written to: {output_dir.resolve()}")
        return 0

    annotation1 = output_dir / "kvq_labels_4k.txt"
    annotation2 = output_dir / "kvq_labels_vdpve.txt"
    create_kvq_annotation(annotation1, dataset1_rows)
    create_kvq_annotation(annotation2, dataset2_rows)

    config_path = Path(args.config)
    checkpoint_path = Path(args.checkpoint)
    if not config_path.is_file():
        raise FileNotFoundError(f"KVQ config does not exist: {config_path}")
    if not checkpoint_path.is_file():
        raise FileNotFoundError(f"KVQ checkpoint does not exist: {checkpoint_path}")
    config = yaml.safe_load(config_path.read_text(encoding="utf-8"))

    device = torch.device(args.device)
    if device.type == "cuda" and not torch.cuda.is_available():
        raise RuntimeError("CUDA was requested but torch.cuda.is_available() is False.")
    model = load_model(config, checkpoint_path, device)

    all_rows = []
    all_rows.extend(
        evaluate_dataset(
            "4K_HEVC",
            dataset1_root,
            dataset1_rows,
            annotation1,
            config,
            model,
            device,
            args.num_workers,
            not args.no_amp,
        )
    )
    # Save a checkpoint after each dataset so a later failure does not lose work.
    fields, normalized = rectangularize(all_rows, ("dataset", "filename", "mos"))
    write_dict_csv(output_dir / "per_video.partial.csv", normalized, fields)

    all_rows.extend(
        evaluate_dataset(
            "VDPVE_train",
            dataset2_root,
            dataset2_rows,
            annotation2,
            config,
            model,
            device,
            args.num_workers,
            not args.no_amp,
        )
    )

    per_video_fields, per_video_rows = rectangularize(
        all_rows, ("dataset", "filename", "mos")
    )
    metrics = compute_metrics(all_rows)
    metrics_fields, metrics_rows = rectangularize(
        metrics,
        (
            "dataset",
            "branch",
            "n",
            "SRCC",
            "KRCC",
            "PLCC",
            "PLCC_logistic",
            "PLCC_logtistic",
        ),
    )
    write_dict_csv(output_dir / "per_video.csv", per_video_rows, per_video_fields)
    write_dict_csv(output_dir / "metrics.csv", metrics_rows, metrics_fields)
    xlsx_written = write_xlsx(output_dir / "kvq_results.xlsx", all_rows, metrics)

    print(f"Results written to: {output_dir.resolve()}")
    print(f"Per-video table: {output_dir / 'per_video.csv'}")
    print(f"Metric table: {output_dir / 'metrics.csv'}")
    if xlsx_written:
        print(f"Excel workbook: {output_dir / 'kvq_results.xlsx'}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("Interrupted.", file=sys.stderr)
        raise SystemExit(130)
